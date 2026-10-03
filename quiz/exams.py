import csv
import io
import random
import re
import zipfile

from openpyxl import Workbook, load_workbook

from .models import Choice, Question
from .roles import STUDENT_RE
from .util import to_int

MAX_ROWS = 500          # questions per upload
MAX_STUDENTS = 20000    # students in one class list (typed, uploaded, or added later in one go)
TOO_MANY_STUDENTS = 'Too many students: a class list can have at most {:,}.'
MAX_BYTES = 1_000_000
MAX_UNPACKED_BYTES = 20_000_000  # an .xlsx is a zip: a tiny file can unpack into gigabytes
MAX_COLUMNS = 30
LETTERS = 'ABCD'

QUESTION_ROWS = [
    ['question', 'option_a', 'option_b', 'option_c', 'option_d', 'correct', 'points'],
    ['What does len("abc") return?', 2, 3, 4, '', 'B', 1],
    ['Which keyword defines a function in Python?', 'func', 'def', 'function', 'lambda', 'B', 2],
]
STUDENT_ROWS = [['register_number'], ['AD3919'], ['AB1234']]
_HEADER_CELLS = {'register_number', 'register number', 'regno', 'reg no', 'email', 'id'}
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():  # Excel stores 1 as 1.0
        value = int(value)
    return str(value).strip()


def read_table(raw, filename='', max_rows=MAX_ROWS + 2):
    """Uploaded CSV (UTF-8) or .xlsx bytes -> list of rows of stripped strings. Raises ValueError with a teacher-friendly message."""
    if len(raw) > MAX_BYTES:
        raise ValueError('File is too large (max 1 MB).')
    name = filename.lower()
    if name.endswith('.xls'):
        raise ValueError('Old .xls files are not supported. Save the file as .xlsx or CSV UTF-8.')
    rows = []
    if name.endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                unpacked = sum(info.file_size for info in archive.infolist())
        except zipfile.BadZipFile:
            raise ValueError('Could not read the Excel file. Save it again as .xlsx or CSV UTF-8.')
        if unpacked > MAX_UNPACKED_BYTES:
            raise ValueError('The Excel file is too large once opened (max 20 MB). Delete unused rows and columns and try again.')
        try:
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for row in book.worksheets[0].iter_rows(max_col=MAX_COLUMNS, values_only=True):  # first sheet, first columns
                rows.append([_cell(c) for c in row])
                if len(rows) >= max_rows:
                    break
            book.close()
        except Exception:
            raise ValueError('Could not read the Excel file. Save it again as .xlsx or CSV UTF-8.')
        return rows
    try:
        text = raw.decode('utf-8-sig')
        for row in csv.reader(io.StringIO(text)):
            rows.append([c.strip() for c in row])
            if len(rows) >= max_rows:
                break
    except (UnicodeDecodeError, csv.Error):
        raise ValueError('File must be UTF-8 CSV or .xlsx (in Excel: Save As, "CSV UTF-8").')
    return rows


def read_upload(f):
    """The bytes of an uploaded file, refusing big ones BEFORE reading them into memory."""
    if f.size > MAX_BYTES:
        raise ValueError('File is too large (max 1 MB).')
    return f.read()


def parse_allowed(text):
    """Class list text -> (unique lowercase emails, bad tokens). 'AD3919' becomes 'ad3919@srmist.edu.in'."""
    emails, bad = [], []
    for token in re.split(r'[\s,;]+', text.strip()):
        if not token:
            continue
        email = token.lower()
        if '@' not in email:
            email += '@srmist.edu.in'
        if STUDENT_RE.match(email):
            emails.append(email)
        else:
            bad.append(token)
    emails = list(dict.fromkeys(emails))
    if len(emails) > MAX_STUDENTS:
        raise ValueError(TOO_MANY_STUDENTS.format(MAX_STUDENTS))
    return emails, bad


def read_student_list(raw, filename=''):
    """Student-list file -> text for parse_allowed. Uses the first column; skips a header row."""
    cells = [row[0] for row in read_table(raw, filename, MAX_STUDENTS + 2) if row and row[0]]  # +2: header, one spare
    if cells and cells[0].lower() in _HEADER_CELLS:
        cells = cells[1:]
    if len(cells) > MAX_STUDENTS:  # say so, rather than silently dropping the rest and turning those students away
        raise ValueError(TOO_MANY_STUDENTS.format(MAX_STUDENTS))
    return '\n'.join(cells)


def parse_questions(raw, filename=''):
    """Question file -> (rows, errors). Columns: question, option_a..option_d, correct (A-D), points (optional).

    Each row is (text, [(option_text, is_correct), ...], points). Blank options are dropped.
    """
    try:
        table = read_table(raw, filename)
    except ValueError as e:
        return [], [str(e)]
    if not table:
        return [], ['The file is empty.']
    header = [h.lower() for h in table[0]]
    missing = {'question', 'option_a', 'option_b', 'correct'} - set(header)
    if missing:
        return [], [f'Missing column(s): {", ".join(sorted(missing))}.']

    rows, errors = [], []
    for n, cells in enumerate(table[1:], start=2):  # row 1 is the header
        if not any(cells):
            continue
        if n - 1 > MAX_ROWS:
            errors.append(f'Too many rows (max {MAX_ROWS}).')
            break
        rec = dict(zip(header, cells))
        get = lambda k: rec.get(k, '')
        question = get('question')
        if not question:
            errors.append(f'Row {n}: question is empty.')
            continue
        if len(question) > 500:
            errors.append(f'Row {n}: question is longer than 500 characters.')
            continue
        columns = [get(f'option_{c.lower()}') for c in LETTERS]
        correct = get('correct').upper()
        if correct not in LETTERS or not columns[LETTERS.index(correct)]:
            errors.append(f'Row {n}: "correct" must be A-D and point to a filled option.')
            continue
        options = [(t, LETTERS[i] == correct) for i, t in enumerate(columns) if t]
        if len(options) < 2:
            errors.append(f'Row {n}: needs at least two options.')
            continue
        if any(len(t) > 300 for t, _ in options):
            errors.append(f'Row {n}: an option is longer than 300 characters.')
            continue
        points = get('points') or '1'
        if to_int(points, 1000) is None or to_int(points, 1000) < 1:
            errors.append(f'Row {n}: points must be a whole number from 1 to 1000.')
            continue
        rows.append((question, options, to_int(points)))
    if not rows and not errors:
        errors.append('No questions found in the file.')
    return rows, errors


def create_questions(owner, rows, question_set=None):
    for text, options, points in rows:
        q = Question.objects.create(text=text, owner=owner, points=points, question_set=question_set)
        # students see the option ids, so they must not follow the file's A-D order (a skewed key would leak)
        Choice.objects.bulk_create(Choice(question=q, text=t, is_correct=ok) for t, ok in random.sample(options, len(options)))
    return len(rows)


def template_bytes(rows, fmt):
    """Template rows -> (bytes, content_type, extension) as CSV (with BOM so Excel reads UTF-8) or .xlsx."""
    if fmt == 'xlsx':
        book = Workbook()
        for row in rows:
            book.active.append(row)
        out = io.BytesIO()
        book.save(out)
        return out.getvalue(), XLSX_TYPE, 'xlsx'
    out = io.StringIO()
    csv.writer(out, lineterminator='\n').writerows(rows)
    return ('﻿' + out.getvalue()).encode('utf-8'), 'text/csv; charset=utf-8', 'csv'
