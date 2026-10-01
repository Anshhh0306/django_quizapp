import io
from datetime import timedelta

from django.contrib.auth.models import Group, User
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from quiz.exam_results import result_rows, result_summary
from quiz.models import Exam, ExamAllowed, ExamAttempt
from quiz.test_exam_take import PASSWORD, TakeBase


class ResultsBase(TakeBase):
    """ab1000 scores well, ab1001 scores low, ab1002 only joined the lobby. A fourth listed student never came."""

    def setUp(self):
        super().setUp()
        self.run_now()
        self.p0 = self.answer_all(self.students[0], 10)  # all 10 right = 11 points
        self.client.post(self.url('exam_submit'))
        self.p1 = self.answer_all(self.students[1], 3)   # 3 right, 7 wrong
        self.client.post(self.url('exam_submit'))
        self.ghost = 'cd5678@srmist.edu.in'
        ExamAllowed.objects.create(exam=self.exam, email=self.ghost)
        for s in self.students:
            ExamAllowed.objects.create(exam=self.exam, email=s.email)

    def rows(self):
        return {r['name']: r for r in result_rows(self.exam)}

    def results_page(self, **params):
        self.login(self.teacher)
        return self.client.get(reverse('exam_results', args=[self.exam.pk]), params)


class ResultRowsTests(ResultsBase):
    def test_counts_for_each_student(self):
        rows = self.rows()
        top, low = rows['ab1000'], rows['ab1001']
        self.assertEqual((top['status'], top['right'], top['wrong'], top['unanswered'], top['score'], top['total'], top['percent']),
                         ('Submitted', 10, 0, 0, 11, 11, 100))
        self.assertEqual((low['right'], low['wrong'], low['unanswered'], low['score'], low['total']),
                         (3, 7, 0, self.worth(self.p1, 3), 11))

    def test_unanswered_questions_are_counted(self):
        # a student who answers 4 of 10 and submits
        other = User.objects.create_user('ab2000', 'ab2000@srmist.edu.in', PASSWORD)
        ExamAllowed.objects.create(exam=self.exam, email=other.email)  # the class list is in force
        self.login(other)
        self.client.post(self.url('exam_consent'), {'agree': 'on'})
        p = self.payload(other)
        for q in p['questions'][:4]:
            self.answer(q['id'], self.choice(q['id'], True))
        self.client.post(self.url('exam_submit'))
        r = self.rows()['ab2000']
        self.assertEqual((r['answered'], r['right'], r['wrong'], r['unanswered']), (4, 4, 0, 6))

    def test_statuses_and_order(self):
        names = [r['name'] for r in result_rows(self.exam)]
        self.assertEqual(names, ['ab1000', 'ab1001', 'ab1002', 'cd5678'])  # best score first, then never-started, then absent
        rows = self.rows()
        self.assertEqual((rows['ab1002']['status'], rows['cd5678']['status']), ('Joined, did not start', 'Did not join'))

    def test_in_progress_student_is_listed_while_the_exam_runs(self):
        self.payload(self.students[2])  # opens the questions but never submits
        self.assertEqual(self.rows()['ab1002']['status'], 'In progress')

    def test_walk_aways_are_submitted_when_results_are_opened_after_time_is_up(self):
        self.payload(self.students[2])
        self.answer(self.exam.questions.first().pk, 1)  # junk id, ignored
        now = timezone.now() - timedelta(minutes=5)
        Exam.objects.filter(pk=self.exam.pk).update(ends_at=now)
        ExamAttempt.objects.filter(exam=self.exam).update(ends_at=now)
        self.assertEqual(self.rows()['ab1002']['status'], 'Submitted')

    def test_summary(self):
        s = result_summary(result_rows(self.exam))
        self.assertEqual((s['class_size'], s['submitted'], s['not_started'], s['absent']), (4, 2, 1, 1))
        self.assertEqual((s['highest'], s['average']), (100, round((100 + self.rows()['ab1001']['percent']) / 2)))


class ResultsPageTests(ResultsBase):
    def test_teacher_sees_the_table(self):
        r = self.results_page()
        self.assertContains(r, 'ab1000')
        self.assertContains(r, '11 / 11')
        self.assertContains(r, 'Did not join')
        self.assertContains(r, 'still running')  # live note while the exam has not ended

    def test_no_live_note_after_the_end(self):
        self.teacher_post('end')
        self.assertNotContains(self.results_page(), 'still running')

    def test_only_the_owner_can_see_results(self):
        other = User.objects.create_user('meena', 'meena@srmist.edu.in', PASSWORD)
        other.groups.add(Group.objects.get(name='Teachers'))
        self.login(other)
        url = reverse('exam_results', args=[self.exam.pk])
        self.assertEqual(self.client.get(url).status_code, 404)
        self.login(self.students[0])
        self.assertRedirects(self.client.get(url), reverse('home'))

    def test_csv_download(self):
        r = self.results_page(format='csv')
        self.assertIn('attachment', r['Content-Disposition'])
        lines = r.content.decode('utf-8-sig').strip().splitlines()
        self.assertTrue(lines[0].startswith('Student,Email,Status'))
        self.assertTrue(lines[1].startswith('ab1000,ab1000@srmist.edu.in,Submitted,10,10,10,0,0,11,11,100'))
        self.assertEqual(len(lines), 5)  # header + 4 students

    def test_excel_download_opens_and_matches(self):
        r = self.results_page(format='xlsx')
        sheet = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in sheet[1]][:3], ['Student', 'Email', 'Status'])
        self.assertEqual(sheet['A2'].value, 'ab1000')
        self.assertEqual(sheet.max_row, 5)

    def test_detail_shows_choices_next_to_correct_answers(self):
        attempt = ExamAttempt.objects.get(user=self.students[1])
        self.login(self.teacher)
        r = self.client.get(reverse('exam_result_detail', args=[self.exam.pk, attempt.pk]))
        items = r.context['items']
        self.assertEqual(len(items), 10)
        self.assertEqual(sum(i['state'] == 'correct' for i in items), 3)
        self.assertEqual(sum(i['state'] == 'wrong' for i in items), 7)
        self.assertContains(r, 'Correct answer:')

    def test_detail_is_owner_only_and_scoped_to_the_exam(self):
        attempt = ExamAttempt.objects.get(user=self.students[1])
        other = User.objects.create_user('meena', 'meena@srmist.edu.in', PASSWORD)
        other.groups.add(Group.objects.get(name='Teachers'))
        self.login(other)
        self.assertEqual(self.client.get(reverse('exam_result_detail', args=[self.exam.pk, attempt.pk])).status_code, 404)
        self.login(self.teacher)
        elsewhere = Exam.objects.create(owner=self.teacher, title='Other', seat_limit=5)
        self.assertEqual(self.client.get(reverse('exam_result_detail', args=[elsewhere.pk, attempt.pk])).status_code, 404)


class LockAfterEndTests(TakeBase):
    def post_detail(self, **data):
        self.login(self.teacher)
        return self.client.post(reverse('exam_detail', args=[self.exam.pk]), data)

    def test_class_list_and_seats_can_change_before_the_end(self):
        self.post_detail(action='allow', students='AB5555')
        self.assertEqual(self.exam.allowed.count(), 1)

    def test_class_list_is_locked_once_the_teacher_ends_the_exam(self):
        self.run_now()
        self.teacher_post('end')
        self.post_detail(action='allow', students='AB5555')
        self.assertEqual(self.exam.allowed.count(), 0)

    def test_class_list_is_locked_once_the_time_is_up(self):
        self.run_now(minutes=-1)
        self.post_detail(action='allow', students='AB5555')
        self.assertEqual(self.exam.allowed.count(), 0)

    def test_seat_limit_is_locked_after_the_end(self):
        self.run_now()
        self.teacher_post('end')
        self.post_detail(action='seats', seat_limit=99)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.seat_limit, 10)

    def test_forms_are_hidden_after_the_end(self):
        self.run_now()
        self.teacher_post('end')
        r = self.client.get(reverse('exam_detail', args=[self.exam.pk]))
        self.assertContains(r, 'The class list and seats are locked')
        self.assertNotContains(r, 'Add students')


class MyExamsTests(ResultsBase):
    def home(self, student):
        self.login(student)
        return self.client.get(reverse('home')).context['my_exams']

    def test_submitted_scheduled_exam_hides_the_score_until_the_end(self):
        card = self.home(self.students[0])[0]
        self.assertEqual((card['state'], card['score'], card['waiting_for_scores']), ('Submitted', None, True))

    def test_score_appears_on_home_after_the_exam_ends(self):
        self.teacher_post('end')
        cards = {c['title']: c for c in self.home(self.students[0])}
        self.assertEqual((cards['Unit 1']['score'], cards['Unit 1']['total']), (11, 11))
        self.assertContains(self.client.get(reverse('home')), 'Score: 11 / 11')

    def test_in_progress_exam_links_back_to_the_questions(self):
        card = self.home(self.students[2])[0]
        self.assertEqual(card['state'], 'Open: start now')
        self.assertTrue(card['url'].endswith('/take/'))

    def test_lobby_exam_links_to_the_lobby(self):
        Exam.objects.filter(pk=self.exam.pk).update(status=Exam.LOBBY, starts_at=None, ends_at=None)
        self.assertTrue(self.home(self.students[2])[0]['url'].endswith('/lobby/'))

    def test_walk_away_is_scored_when_the_student_opens_home(self):
        self.payload(self.students[2])
        past = timezone.now() - timedelta(minutes=5)
        Exam.objects.filter(pk=self.exam.pk).update(ends_at=past)
        ExamAttempt.objects.filter(exam=self.exam).update(ends_at=past)
        card = self.home(self.students[2])[0]
        self.assertEqual(card['state'], 'Submitted')
        self.assertEqual(card['score'], 0)  # the exam is over, so the score is visible

    def test_students_only_see_their_own_exams(self):
        outsider = User.objects.create_user('cd7777', 'cd7777@srmist.edu.in', PASSWORD)
        self.assertEqual(self.home(outsider), [])

    def test_open_exam_shows_the_score_straight_away(self):
        Exam.objects.filter(pk=self.exam.pk).update(mode=Exam.OPEN)
        self.assertEqual(self.home(self.students[0])[0]['score'], 11)
